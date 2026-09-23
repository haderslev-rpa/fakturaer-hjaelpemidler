"""Læs leverandører og kontostrenge fra SharePoint-Excel i memory."""
from __future__ import annotations

from functools import lru_cache
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook
from q_sharepoint_api.sp_api import get_client

from konfiguration import (
    SHAREPOINT_LEVERANDOER_FILE_PATH,
    SHAREPOINT_SITE_NAME,
)


def _normaliser(value: Any) -> str:
    return " ".join(str(value or "").strip().casefold().split())


@lru_cache(maxsize=1)
def hent_excelgrundlag() -> dict[str, list[dict[str, Any]]]:
    """Download Excel én gang pr. proces og returnér de tre faner som rækker."""
    client = get_client()
    site_id = client.get_site_id(SHAREPOINT_SITE_NAME)
    result = client.download_file_to_memory_by_path(
        site_id=site_id,
        file_path=SHAREPOINT_LEVERANDOER_FILE_PATH,
        save_dir=None,
    )
    file_bytes = result.get("file_bytes")
    if not file_bytes:
        raise ValueError("SharePoint-filen blev hentet, men file_bytes er tom.")

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    return {
        "leverandoer": _laes_ark(workbook, ("leverandoer", "leverandør")),
        "kontostrenge_2026": _laes_ark(
            workbook,
            ("kontostrenge 2026", "kontostreng 2026"),
        ),
        "kontostrenge_andre_aar": _laes_ark(
            workbook,
            ("kontostrenge andre år", "kontostreng andre år"),
        ),
    }


def _laes_ark(workbook, aliases: tuple[str, ...]) -> list[dict[str, Any]]:
    sheet = None
    wanted = {_normaliser(alias) for alias in aliases}
    for sheet_name in workbook.sheetnames:
        if _normaliser(sheet_name) in wanted:
            sheet = workbook[sheet_name]
            break
    if sheet is None:
        raise ValueError(
            f"Excel-filen mangler en arkfane med et af navnene: {aliases!r}. "
            f"Fundne faner: {workbook.sheetnames!r}."
        )

    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(rows)]
    except StopIteration as error:
        raise ValueError(f"Arkfanen {sheet.title!r} er tom.") from error

    result = []
    for values in rows:
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
            if headers[index]
        }
        if any(value not in (None, "") for value in row.values()):
            result.append(row)
    return result


def hent_tilladte_leverandoerer() -> list[dict[str, str]]:
    """Returnér leverandørfanens Fakturaafsender og CVR nr."""
    rows = hent_excelgrundlag()["leverandoer"]
    suppliers = []
    seen = set()
    for row in rows:
        name = str(row.get("Fakturaafsender") or "").strip()
        cvr = normaliser_cvr(row.get("CVR nr"))
        if not name and not cvr:
            continue
        if not name or not cvr:
            raise ValueError(
                "En række i fanen 'leverandoer' mangler Fakturaafsender eller CVR nr."
            )
        key = (name.casefold(), cvr)
        if key not in seen:
            seen.add(key)
            suppliers.append({"fakturaafsender": name, "cvr_nr": cvr})
    if not suppliers:
        raise ValueError("Fanen 'leverandoer' indeholder ingen gyldige leverandører.")
    return suppliers


def hent_kontostrengsraekker(arknoegle: str) -> list[dict[str, Any]]:
    """Returnér rækker fra den valgte kontostrengsfane."""
    rows = hent_excelgrundlag().get(arknoegle)
    if rows is None:
        raise ValueError(f"Ukendt kontostrengsark: {arknoegle!r}.")
    return rows


def normaliser_cvr(value: Any) -> str:
    """Returnér CVR som cifre uden DK-prefix."""
    text = re.sub(r"\D", "", str(value or ""))
    return text[-8:] if len(text) >= 8 else text
